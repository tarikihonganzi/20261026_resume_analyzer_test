import asyncio
import io
import time
import zipfile
from collections.abc import Callable

import pytest
from openpyxl import Workbook

from app.domain.errors import (
    ConversionFailedError,
    EmptyResumeError,
    FileTooLargeError,
    InvalidResumeInputError,
    ResumeError,
    UnsupportedFormatError,
)
from app.ingest import resume
from app.ingest.resume import ResumeInput, check_of, load_resume
from fixtures import resume_factory as factory

MB = 1024 * 1024


async def load_file(filename: str, content: bytes, *, max_size_mb: int = 0) -> ResumeInput:
    return await load_resume(
        filename=filename, content=content, pasted_text=None, max_size_mb=max_size_mb
    )


async def load_text(text: str, *, max_size_mb: int = 0) -> ResumeInput:
    return await load_resume(filename=None, content=None, pasted_text=text, max_size_mb=max_size_mb)


# --- 変換できる形式 ---


@pytest.mark.parametrize(
    ("filename", "content", "format_"),
    [
        pytest.param("田中_職務経歴書.docx", factory.make_docx(), "docx", id="docx"),
        pytest.param("経歴.xlsx", factory.make_xlsx(), "xlsx", id="xlsx"),
        pytest.param("経歴.pptx", factory.make_pptx(), "pptx", id="pptx"),
        pytest.param("経歴.txt", factory.make_text().encode("utf-8"), "txt", id="txt"),
        pytest.param("経歴.md", factory.make_text().encode("utf-8"), "md", id="md"),
    ],
)
async def test_office_and_text_files_are_converted_to_text(
    filename: str, content: bytes, format_: str
) -> None:
    result = await load_file(filename, content)

    assert result.format == format_
    assert result.source_name == filename
    assert result.pdf is None
    assert result.text is not None
    assert "1級建築施工管理技士" in result.text


async def test_shift_jis_text_file_is_decoded() -> None:
    result = await load_file("経歴.txt", factory.make_text().encode("cp932"))

    assert result.text is not None
    assert "1級建築施工管理技士" in result.text


async def test_file_without_extension_is_guessed_by_content() -> None:
    result = await load_file("resume", factory.make_text().encode("utf-8"))

    assert result.text is not None
    assert "1級建築施工管理技士" in result.text


# --- PDF ---


@pytest.mark.parametrize(
    "make",
    [
        factory.make_pdf_digital,
        pytest.param(
            factory.make_pdf_scanned,
            marks=pytest.mark.skipif(
                factory.find_japanese_font() is None, reason="日本語フォントが見つからない"
            ),
        ),
    ],
)
async def test_pdf_is_kept_as_bytes_and_not_converted(make: Callable[[], bytes]) -> None:
    content = make()

    result = await load_file("職歴書.pdf", content)

    assert result.format == "pdf"
    assert result.pdf == content
    assert result.text is None


async def test_extension_is_case_insensitive() -> None:
    result = await load_file("RESUME.PDF", factory.make_pdf_digital())

    assert result.format == "pdf"


async def test_pdf_is_detected_by_its_header_even_with_another_extension() -> None:
    result = await load_file("resume.txt", factory.make_pdf_digital())

    assert result.format == "pdf"
    assert result.pdf is not None


async def test_file_named_pdf_without_a_pdf_header_is_an_error() -> None:
    with pytest.raises(ConversionFailedError):
        await load_file("resume.pdf", b"this is not a pdf" * 10)


# --- 貼り付け ---


async def test_pasted_text_uses_the_fixed_source_name() -> None:
    result = await load_text(factory.make_text())

    assert result.format == "text"
    assert result.source_name == "貼り付けテキスト"
    assert result.pdf is None
    assert result.text == factory.make_text()


@pytest.mark.parametrize("text", ["", "   ", "\n\n\t", "短い", "あ" * 19])
async def test_empty_or_extremely_short_pasted_text_is_an_error(text: str) -> None:
    with pytest.raises(EmptyResumeError) as excinfo:
        await load_text(text)

    assert "テキスト化できませんでした" in str(excinfo.value)


async def test_pasted_text_of_the_minimum_length_is_accepted() -> None:
    assert (await load_text("あ" * 20)).text == "あ" * 20


# --- 入力の組み合わせ ---


async def test_both_file_and_text_is_an_error() -> None:
    with pytest.raises(InvalidResumeInputError):
        await load_resume(
            filename="a.txt",
            content=factory.make_text().encode(),
            pasted_text=factory.make_text(),
            max_size_mb=0,
        )


async def test_neither_file_nor_text_is_an_error() -> None:
    with pytest.raises(InvalidResumeInputError):
        await load_resume(filename=None, content=None, pasted_text=None, max_size_mb=0)


async def test_file_content_without_a_name_is_an_error() -> None:
    with pytest.raises(InvalidResumeInputError):
        await load_resume(filename=None, content=b"abc", pasted_text=None, max_size_mb=0)


# --- サイズ上限 ---


async def test_file_over_the_limit_is_rejected_with_the_limit_in_the_message() -> None:
    content = b"a" * (1 * MB + 1)

    with pytest.raises(FileTooLargeError) as excinfo:
        await load_file("resume.txt", content, max_size_mb=1)

    assert excinfo.value.limit_mb == 1
    assert "1 MB" in str(excinfo.value)


async def test_file_of_exactly_the_limit_is_accepted() -> None:
    content = b"a" * (1 * MB)

    result = await load_file("resume.txt", content, max_size_mb=1)

    assert result.text is not None


async def test_zero_means_no_limit() -> None:
    result = await load_file("resume.txt", b"a" * (3 * MB), max_size_mb=0)

    assert result.text is not None


async def test_pdf_is_also_subject_to_the_limit() -> None:
    with pytest.raises(FileTooLargeError):
        await load_file("resume.pdf", factory.make_pdf_digital() + b"0" * (1 * MB), max_size_mb=1)


async def test_pasted_text_is_also_subject_to_the_limit() -> None:
    with pytest.raises(FileTooLargeError):
        await load_text("あ" * (MB // 3 + 1), max_size_mb=1)


# --- 空・短い ---


async def test_docx_without_text_is_an_error() -> None:
    with pytest.raises(EmptyResumeError):
        await load_file("empty.docx", factory.make_empty_docx())


async def test_xlsx_without_cells_is_an_error() -> None:
    buffer = io.BytesIO()
    Workbook().save(buffer)

    with pytest.raises(EmptyResumeError):
        await load_file("empty.xlsx", buffer.getvalue())


async def test_text_file_that_is_too_short_is_an_error() -> None:
    with pytest.raises(EmptyResumeError):
        await load_file("short.txt", "経歴なし".encode())


# --- 読めない・未対応 ---


async def test_broken_docx_is_a_conversion_error_without_leaking_content() -> None:
    content = b"PK\x03\x04SECRET-MARKER-0123456789"

    with pytest.raises(ConversionFailedError) as excinfo:
        await load_file("broken.docx", content)

    assert "SECRET-MARKER" not in str(excinfo.value)


async def test_docx_named_as_text_is_a_conversion_error() -> None:
    # MarkItDown は拡張子を信じて、バイナリを文字として読んでしまう
    with pytest.raises(ConversionFailedError):
        await load_file("resume.txt", factory.make_docx())


async def test_binary_data_named_as_text_is_a_conversion_error() -> None:
    with pytest.raises(ConversionFailedError):
        await load_file("resume.txt", b"abc\x00def" * 100)


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        pytest.param("resume.bin", bytes(range(256)) * 10, id="bin"),
        pytest.param("setup.exe", b"MZ" + bytes(range(256)) * 4, id="exe"),
    ],
)
async def test_unsupported_file_is_an_error_that_names_the_supported_formats(
    filename: str, content: bytes
) -> None:
    with pytest.raises(UnsupportedFormatError) as excinfo:
        await load_file(filename, content)

    assert "PDF" in str(excinfo.value)
    assert "docx" in str(excinfo.value)


async def test_zip_is_not_accepted() -> None:
    # zip は中身を再帰的に展開する変換器があり、zip 爆弾の危険がある
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("resume.txt", factory.make_text())

    with pytest.raises(UnsupportedFormatError):
        await load_file("resume.zip", buffer.getvalue())


# --- Copilot のレビューで指摘された、取り込みの抜け穴 ---


@pytest.mark.parametrize("filename", ["resume", "resume.dat", "resume.zip", "resume.docx.bak"])
async def test_generic_zip_is_rejected_whatever_its_file_name(filename: str) -> None:
    # 拡張子を外す・変えるだけで、zip の再帰展開を行う変換器に回ってはならない
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("a.txt", factory.make_text())

    with pytest.raises(ResumeError):
        await load_file(filename, buffer.getvalue())


def test_zip_converter_is_not_registered_in_markitdown() -> None:
    registrations = resume._markitdown()._converters  # pyright: ignore[reportPrivateUsage]
    registered = {type(r.converter).__name__ for r in registrations}

    assert "ZipConverter" not in registered
    assert {"DocxConverter", "XlsxConverter", "PptxConverter", "PlainTextConverter"} <= registered


async def test_docx_without_an_extension_is_recognized_by_its_content() -> None:
    result = await load_file("resume", factory.make_docx())

    assert result.format == "docx"
    assert result.text is not None
    assert "1級建築施工管理技士" in result.text


@pytest.mark.parametrize(
    "extension", ["txt", "text", "md", "markdown", "jsonl", "json", "csv", "html"]
)
async def test_office_file_with_any_text_extension_is_a_conversion_error(extension: str) -> None:
    with pytest.raises(ConversionFailedError):
        await load_file(f"resume.{extension}", factory.make_docx())


@pytest.mark.parametrize("encoding", ["utf-16", "utf-32", "utf-8-sig"])
async def test_unicode_text_with_a_bom_is_accepted(encoding: str) -> None:
    # UTF-16・UTF-32 は ASCII の文字に NUL バイトを含むが、正常なテキスト
    content = factory.make_text().encode(encoding)

    result = await load_file("経歴.txt", content)

    assert result.text is not None
    assert "1級建築施工管理技士" in result.text


async def test_pdf_header_text_inside_a_text_file_does_not_make_it_a_pdf() -> None:
    memo = "開発メモ: PDF は先頭が %PDF-1.7 で始まる。\n%PDF-1.4 という行もある。\n"
    body = factory.make_text() + memo

    result = await load_file("経歴.txt", body.encode("utf-8"))

    assert result.format == "txt"
    assert result.pdf is None


@pytest.mark.parametrize("prefix", [b"", b"\xef\xbb\xbf", b"\n\n", b"  "])
async def test_pdf_header_is_recognized_at_the_start_of_the_file(prefix: bytes) -> None:
    result = await load_file("resume.dat", prefix + factory.make_pdf_digital())

    assert result.format == "pdf"


async def test_pdf_with_junk_before_the_header_is_accepted_only_with_the_pdf_extension() -> None:
    content = b"junk bytes\n" + factory.make_pdf_digital()

    assert (await load_file("resume.pdf", content)).format == "pdf"
    with pytest.raises(ResumeError):
        await load_file("resume.dat", content)


async def test_pdf_header_without_a_version_is_not_a_pdf() -> None:
    with pytest.raises(ConversionFailedError):
        await load_file("resume.pdf", b"%PDF- but not really a pdf" * 20)


@pytest.mark.parametrize(
    "text",
    [
        "-" * 40,
        "=" * 40,
        "- - -\n" * 7,
        "_ _ _\n" * 7,
        "* * *\n" * 7,
        "| | |\n" * 7,
        "| --- | --- | --- | --- | --- | --- |\n" * 6,
        "|" * 40,
        "# 見出しだけの文書です。長い見出しだけでは職歴書になりません\n## もう一つの長い見出しです",
    ],
    ids=[
        "hyphens",
        "equals",
        "spaced-hyphens",
        "spaced-underscores",
        "spaced-asterisks",
        "empty-table-rows",
        "table-separators",
        "pipes",
        "headings-only",
    ],
)
async def test_text_made_only_of_markup_is_empty(text: str) -> None:
    with pytest.raises(EmptyResumeError):
        await load_text(text)


async def test_markup_around_real_text_is_not_counted_but_the_text_is() -> None:
    text = "# 職歴\n\n| 期間 | 会社 |\n| --- | --- |\n| 2015〜2020 | 株式会社サンプル建設 |\n"

    result = await load_text(text)

    assert result.text == text


@pytest.mark.parametrize("sheet_names", [["あ" * 30], ["シート1", "シート2", "シート3"]])
async def test_workbook_with_only_empty_sheets_is_empty_whatever_the_sheet_names(
    sheet_names: list[str],
) -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = sheet_names[0]
    for name in sheet_names[1:]:
        workbook.create_sheet(name)
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(EmptyResumeError):
        await load_file("empty.xlsx", buffer.getvalue())


# --- 出典名 ---


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("C:\\Users\\taro\\職務経歴書.docx", "職務経歴書.docx"),
        ("../../etc/職歴.txt", "職歴.txt"),
        ("履歴\n書.txt", "履歴書.txt"),
    ],
)
async def test_source_name_is_sanitized(filename: str, expected: str) -> None:
    result = await load_file(filename, factory.make_text().encode("utf-8"))

    assert result.source_name == expected


# --- イベントループを止めない ---


async def test_conversion_does_not_block_the_event_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    def slow_convert(content: bytes, extension: str, filename: str) -> str:
        time.sleep(0.3)
        return factory.make_text()

    monkeypatch.setattr(resume, "_convert", slow_convert)
    ticks = 0

    async def ticker() -> None:
        nonlocal ticks
        while True:
            await asyncio.sleep(0.02)
            ticks += 1

    task = asyncio.create_task(ticker())
    try:
        await load_file("resume.docx", b"PK\x03\x04dummy")
    finally:
        task.cancel()

    assert ticks >= 5


# --- 検証結果（/api/resumes/check の応答の元） ---


async def test_check_of_a_converted_file_reports_the_length() -> None:
    result = await load_file("経歴.docx", factory.make_docx())

    check = check_of(result)

    assert check.format == "docx"
    assert check.source_name == "経歴.docx"
    assert result.text is not None
    assert check.char_count == len(result.text)


async def test_check_of_a_pdf_has_no_length() -> None:
    check = check_of(await load_file("経歴.pdf", factory.make_pdf_digital()))

    assert check.format == "pdf"
    assert check.char_count is None


async def test_check_contains_nothing_but_the_summary() -> None:
    check = check_of(await load_file("経歴.docx", factory.make_docx()))

    assert set(check.model_dump()) == {"format", "source_name", "char_count"}
