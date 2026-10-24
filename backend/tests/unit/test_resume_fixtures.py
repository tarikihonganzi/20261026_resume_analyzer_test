"""架空の職歴書の生成（tests/fixtures/resume_factory.py）が、各形式として正しいことを確認する。"""

import io
import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from fixtures import resume_factory as factory

needs_font = pytest.mark.skipif(
    factory.find_japanese_font() is None, reason="日本語フォントが見つからない"
)


def test_every_company_and_the_name_are_marked_as_fictional() -> None:
    markers = ("架空", "サンプル")

    assert any(m in factory.NAME for m in markers)
    for _, company, _, _ in factory.CAREER:
        assert any(m in company for m in markers), company


def test_text_has_the_common_content() -> None:
    text = factory.make_text()

    assert "1級建築施工管理技士" in text
    assert "株式会社サンプル建設" in text
    assert factory.INJECTION_LINE not in text


def test_injection_line_is_added_only_when_requested() -> None:
    assert factory.INJECTION_LINE in factory.make_text(injection=True)
    assert factory.INJECTION_LINE not in factory.make_text()


def test_docx_can_be_read_back() -> None:
    doc = Document(io.BytesIO(factory.make_docx()))

    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert "1級建築施工管理技士" in paragraphs
    assert "架空工務店株式会社" in cells


def test_xlsx_has_one_sheet_per_topic() -> None:
    wb = load_workbook(io.BytesIO(factory.make_xlsx()))

    assert wb.sheetnames == ["基本情報", "保有資格", "職歴", "スキル"]
    values = [cell for row in wb["職歴"].iter_rows(values_only=True) for cell in row]
    assert "株式会社サンプル建設" in values


def test_pptx_has_one_slide_per_topic() -> None:
    data = factory.make_pptx()

    assert len(Presentation(io.BytesIO(data)).slides) == 5
    assert any("Revit" in text for text in factory.pptx_texts(data))


def test_digital_pdf_has_text_layer_fonts() -> None:
    data = factory.make_pdf_digital()

    assert data.startswith(b"%PDF-")
    assert b"/Font" in data


def test_digital_pdf_is_reproducible() -> None:
    assert factory.make_pdf_digital() == factory.make_pdf_digital()


@needs_font
def test_scanned_pdf_has_only_an_image() -> None:
    data = factory.make_pdf_scanned()

    assert data.startswith(b"%PDF-")
    assert b"/Image" in data
    assert b"/Font" not in data


@needs_font
def test_scanned_pdf_is_larger_than_digital_pdf() -> None:
    # 画像を含むので、テキストだけの PDF よりずっと大きい（取り違えの検出）
    assert len(factory.make_pdf_scanned()) > len(factory.make_pdf_digital())


def test_broken_docx_cannot_be_opened() -> None:
    data = factory.make_broken_docx()

    assert data.startswith(b"PK")
    with pytest.raises(Exception):  # noqa: B017 - zipfile.BadZipFile など、形式の破損を表す例外
        Document(io.BytesIO(data))


def test_empty_docx_has_no_text() -> None:
    doc = Document(io.BytesIO(factory.make_empty_docx()))

    assert "".join(p.text for p in doc.paragraphs).strip() == ""


def test_all_fixtures_have_unique_non_empty_files() -> None:
    fixtures = factory.all_fixtures()

    names = [f.filename for f in fixtures]
    assert len(names) == len(set(names))
    assert all(f.content for f in fixtures)
    assert {"resume.txt", "resume.docx", "resume.xlsx", "resume.pptx", "resume_digital.pdf"} <= set(
        names
    )
    assert {"broken.docx", "empty.docx"} <= set(names)


def test_command_line_writes_every_fixture(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(sys, "argv", ["resume_factory.py", str(tmp_path)])

    factory.main()

    written = {p.name for p in tmp_path.iterdir()}
    assert written == {f.filename for f in factory.all_fixtures()}
